import io
import json
import math
import re
from collections.abc import Sequence
from dataclasses import dataclass

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import KnowledgeChunk, KnowledgeDocument
from app.services.gemini import embed_texts

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_CHUNKS_PER_DOCUMENT = 200
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 180
EMBEDDING_BATCH_SIZE = 32
RETRIEVAL_LIMIT = 5
MINIMUM_RELEVANCE = 0.25
SUPPORTED_TYPES = {".txt", ".md", ".pdf", ".docx"}


@dataclass
class TextChunk:
    content: str
    page_number: int | None


@dataclass
class RetrievedChunk:
    filename: str
    page_number: int | None
    content: str
    score: float

    @property
    def citation(self) -> str:
        suffix = f", halaman {self.page_number}" if self.page_number else ""
        return f"{self.filename}{suffix}"


def _split_text(text: str, page_number: int | None) -> list[TextChunk]:
    normalized = re.sub(r"\s+", " ", text).strip()
    chunks: list[TextChunk] = []
    start = 0
    while start < len(normalized):
        end = min(start + CHUNK_SIZE, len(normalized))
        if end < len(normalized):
            boundary = normalized.rfind(" ", start + CHUNK_SIZE // 2, end)
            if boundary > start:
                end = boundary
        content = normalized[start:end].strip()
        if content:
            chunks.append(TextChunk(content=content, page_number=page_number))
        if end >= len(normalized):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return chunks


async def _read_upload(upload: UploadFile) -> tuple[str, str, list[TextChunk]]:
    filename = (upload.filename or "").replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not filename or len(filename) > 255:
        raise HTTPException(status_code=400, detail="Nama file tidak valid.")
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in SUPPORTED_TYPES:
        raise HTTPException(
            status_code=415,
            detail="Format dokumen yang didukung: PDF, TXT, MD, dan DOCX.",
        )
    data = await upload.read(MAX_FILE_BYTES + 1)
    if not data:
        raise HTTPException(status_code=400, detail="File dokumen kosong.")
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="Ukuran file maksimum adalah 5 MB.")

    content_type = upload.content_type or "application/octet-stream"
    try:
        if extension in {".txt", ".md"}:
            text = data.decode("utf-8-sig")
            chunks = _split_text(text, None)
        elif extension == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            if len(reader.pages) > 300:
                raise HTTPException(status_code=413, detail="PDF tidak boleh melebihi 300 halaman.")
            chunks = [
                chunk
                for page_number, page in enumerate(reader.pages, start=1)
                for chunk in _split_text(page.extract_text() or "", page_number)
            ]
        else:
            document = DocxDocument(io.BytesIO(data))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            chunks = _split_text(text, None)
    except (
        UnicodeDecodeError,
        OSError,
        ValueError,
        PdfReadError,
        PackageNotFoundError,
    ) as exc:
        raise HTTPException(
            status_code=422,
            detail="File tidak dapat dibaca. Pastikan format dan isi dokumen valid.",
        ) from exc

    if not chunks:
        raise HTTPException(status_code=422, detail="Tidak ada teks yang dapat diindeks.")
    if len(chunks) > MAX_CHUNKS_PER_DOCUMENT:
        raise HTTPException(
            status_code=413,
            detail="Dokumen terlalu panjang untuk satu unggahan. Bagi dokumen lalu coba lagi.",
        )
    return filename, content_type, chunks


def _batches(items: Sequence[TextChunk], batch_size: int) -> list[Sequence[TextChunk]]:
    return [items[index : index + batch_size] for index in range(0, len(items), batch_size)]


async def index_upload(db: Session, upload: UploadFile) -> KnowledgeDocument:
    filename, content_type, chunks = await _read_upload(upload)
    embeddings: list[list[float]] = []
    for batch in _batches(chunks, EMBEDDING_BATCH_SIZE):
        embeddings.extend(
            embed_texts([chunk.content for chunk in batch], task_type="RETRIEVAL_DOCUMENT")
        )

    document = KnowledgeDocument(
        filename=filename,
        content_type=content_type,
        chunk_count=len(chunks),
    )
    document.chunks = [
        KnowledgeChunk(
            content=chunk.content,
            page_number=chunk.page_number,
            embedding=json.dumps(vector, separators=(",", ":")),
        )
        for chunk, vector in zip(chunks, embeddings, strict=True)
    ]
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


def _cosine_similarity(left: Sequence[float], right: Sequence[float]) -> float:
    if len(left) != len(right):
        return 0.0
    left_norm = math.sqrt(sum(value * value for value in left))
    right_norm = math.sqrt(sum(value * value for value in right))
    if left_norm == 0 or right_norm == 0:
        return 0.0
    return sum(a * b for a, b in zip(left, right, strict=True)) / (left_norm * right_norm)


def retrieve_context(db: Session, query: str) -> list[RetrievedChunk]:
    stored_chunks = db.execute(
        select(KnowledgeChunk, KnowledgeDocument.filename)
        .join(KnowledgeDocument, KnowledgeChunk.document_id == KnowledgeDocument.id)
    ).all()
    if not stored_chunks:
        return []

    query_vector = embed_texts([query], task_type="RETRIEVAL_QUERY")[0]
    scored = [
        RetrievedChunk(
            filename=filename,
            page_number=chunk.page_number,
            content=chunk.content,
            score=_cosine_similarity(query_vector, json.loads(chunk.embedding)),
        )
        for chunk, filename in stored_chunks
    ]
    return [
        item
        for item in sorted(scored, key=lambda item: item.score, reverse=True)
        if item.score >= MINIMUM_RELEVANCE
    ][:RETRIEVAL_LIMIT]


def remove_document(db: Session, document_id: int) -> bool:
    document = db.get(KnowledgeDocument, document_id)
    if document is None:
        return False
    db.delete(document)
    db.commit()
    return True
